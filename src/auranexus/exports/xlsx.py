"""Excel exports (TestStarter tracker workbook)."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

STATUSES = ["Not Started", "Ready", "In Progress", "Blocked", "Passed", "Failed", "Retest"]
COLUMNS = [
    ("Test ID", "test_id"), ("Scenario", "scenario"), ("Process area", "process_area"), ("Scope item", "scope_item"),
    ("Requirement IDs", "requirement_ids"), ("Step owner", "step_owner"), ("Role", "role"), ("System / client", "system_client"),
    ("Test data ref", "test_data_ref"), ("Planned start", "planned_start"), ("Planned end", "planned_end"),
    ("Actual start", "actual_start"), ("Actual end", "actual_end"), ("Status", "status"),
    ("Predecessors", "predecessor_ids"), ("Successors", "successor_ids"), ("Dependency type", "dependency_type"),
    ("Blocker reason", "blocker_reason"), ("Defect ID", "defect_id"), ("Defect severity", "defect_severity"),
]


def _cell(v: Any) -> Any:
    return ", ".join(map(str, v)) if isinstance(v, list) else v


def test_tracker_workbook(cases: list[dict[str, Any]], roster: list[dict[str, Any]] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    head_fill = PatternFill("solid", fgColor="1F3864")
    for i, (label, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
        ws.column_dimensions[c.column_letter].width = max(14, len(label) + 2)
    succ: dict[str, list[str]] = {}
    for c in cases:
        for p in c.get("predecessor_ids", []):
            succ.setdefault(p, []).append(c["test_id"])
    for r, case in enumerate(sorted(cases, key=lambda x: x["test_id"]), 2):
        row = {**case, "successor_ids": succ.get(case["test_id"], [])}
        for i, (_, key) in enumerate(COLUMNS, 1):
            ws.cell(row=r, column=i, value=_cell(row.get(key)))
    last = max(len(cases) + 1, 2)
    status_col = [k for _, k in COLUMNS].index("status") + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(row=2, column=status_col).coordinate}:{ws.cell(row=last + 500, column=status_col).coordinate}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=last, column=len(COLUMNS)).coordinate}"

    ro = wb.create_sheet("Roster")
    for i, h in enumerate(["Tester", "Backup", "Contact", "Availability"], 1):
        ro.cell(row=1, column=i, value=h).font = Font(bold=True)
    for r, p in enumerate(roster or [], 2):
        for i, k in enumerate(["tester", "backup", "contact", "availability"], 1):
            ro.cell(row=r, column=i, value=p.get(k))

    db = wb.create_sheet("Dashboard")
    col = ws.cell(row=1, column=status_col).column_letter
    db["A1"], db["B1"] = "Status", "Count"
    for i, s in enumerate(STATUSES, 2):
        db.cell(row=i, column=1, value=s)
        db.cell(row=i, column=2, value=f'=COUNTIF(Tracker!${col}$2:${col}$5000,A{i})')
    n = len(STATUSES) + 2
    db.cell(row=n, column=1, value="Total").font = Font(bold=True)
    db.cell(row=n, column=2, value=f"=SUM(B2:B{n - 1})")
    db.cell(row=n + 1, column=1, value="% executed")
    db.cell(row=n + 1, column=2, value=f"=IF(B{n}=0,0,(B6+B7)/B{n})").number_format = "0.0%"
    chart = BarChart()
    chart.title, chart.y_axis.title = "Test cases by status", "Count"
    chart.add_data(Reference(db, min_col=2, min_row=1, max_row=len(STATUSES) + 1), titles_from_data=True)
    chart.set_categories(Reference(db, min_col=1, min_row=2, max_row=len(STATUSES) + 1))
    db.add_chart(chart, "D2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
